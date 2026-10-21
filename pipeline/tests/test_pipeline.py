"""All fixtures here are explicitly synthetic and contain no health estimates."""

import json

import duckdb
import openpyxl
import pandas as pd
import pandera.pandas as pa
import pytest

from visceral_pipeline.aggregate import aggregate
from visceral_pipeline.causes import cause_from_row, hierarchy
from visceral_pipeline.cli import copy_parquet
from visceral_pipeline.constants import AGES, CONFIG, ROOT
from visceral_pipeline.contract import finite_tree, validate_export
from visceral_pipeline.export import export, sig
from visceral_pipeline.ghe import number, parse_reference, parse_workbook
from visceral_pipeline.mapping import load_mapping, validate_mapping
from visceral_pipeline.methods import save_method_comparison
from visceral_pipeline.reference_validation import validate_reference_rates
from visceral_pipeline.standardize import collapse_weights, load_weights, standardize
from visceral_pipeline.validate import reconcile, reviewed_exceptions, schemas


@pytest.fixture
def organs():
    return [
        o["id"] for o in json.loads((ROOT / "contracts/organs.json").read_text(encoding="utf-8"))["organs"]
    ]


@pytest.fixture
def tiny_workbook(tmp_path):
    book = openpyxl.Workbook()
    book.active.title = "Notes"
    for sheet in AGES:
        ws = book.create_sheet(sheet)
        ws.append(["SYNTHETIC test values ('000)"])
        ws.append([None] * 4)
        ws.append(["Sex", "GHE code", None, "GHE cause", None, None, "Country", "Test country"])
        ws.append([None] * 6 + ["ISO-3 Code", "TST"])
        ws.append([None] * 7 + [4])
        for sex, factor in [("Persons", 1), ("Males", 0.6), ("Females", 0.4)]:
            ws.append([sex, None, None, "Population ('000) (2)", None, None, None, 10 * factor])
            ws.append([sex, 0, None, "All Causes", None, None, None, factor])
            ws.append([sex, 10, "I.", "Synthetic condition", None, None, None, factor])
            ws.append([sex, 20, None, "A.", "Synthetic child", None, None, factor])
    path = tmp_path / "synthetic.xlsx"
    book.save(path)
    return path


def test_parser_content_detection_and_units(tiny_workbook):
    result = parse_workbook(tiny_workbook, 2023, "deaths")
    assert len(result.observations) == 8 * 3 * 3
    assert len(result.population) == 24
    row = result.observations.query("sex=='both' and age_group=='all' and cause_code==0").iloc[0]
    assert row.value == 1000
    assert result.population.population.max() == 10000
    assert result.causes.parent.tolist()[1:] == [0, 10]
    assert result.locations.data_quality.tolist() == [4]


def test_parser_rejects_layout_change(tiny_workbook):
    book = openpyxl.load_workbook(tiny_workbook)
    book["All ages"]["B3"] = "Changed header"
    book.save(tiny_workbook)
    with pytest.raises(ValueError):
        parse_workbook(tiny_workbook, 2023, "deaths")


def test_leishmaniasis_nested_printed_label():
    row = ["Persons", 261, None, None, None, "e.1.", "Visceral leishmaniasis"]
    assert cause_from_row(row)["level"] == 5
    with pytest.raises(ValueError):
        hierarchy([{"code": 0, "name": "All", "level": 0}, {"code": 2, "name": "Bad", "level": 3}])


def test_real_mapping_partitions_all_causes(organs):
    mapping = load_mapping(organs)
    validate_mapping(mapping, organs)
    nodes = mapping.set_index("code")
    leaves = set(mapping.code) - set(mapping.parent.dropna())
    # Assign an independent token to every leaf, then sum primary descendant coverage.
    # Every token is counted exactly once, even when the primary is an ancestor.
    coverage = dict.fromkeys(leaves, 0)
    for leaf in leaves:
        cursor = leaf
        while pd.notna(cursor):
            if nodes.loc[cursor, "relation"] == "primary":
                coverage[leaf] += 1
            cursor = nodes.loc[cursor, "parent"]
    assert set(coverage.values()) == {1}
    assert len(mapping) == len(set(mapping.code))
    assert nodes.loc[261, "parent"] == 260
    assert not ((mapping.organ == "muscles") & (mapping.relation == "primary")).any()
    bad = mapping.copy()
    bad.loc[bad.code == 0, ["relation", "organ"]] = ["primary", "heart"]
    with pytest.raises(ValueError, match="Overlapping|primary"):
        validate_mapping(bad, organs)


def test_standardization_hand_example():
    assert standardize({"young": 10.0, "old": 100.0}, {"young": 0.75, "old": 0.25}) == 32.5
    assert standardize({"young": None, "old": 100.0}, {"young": 0.75, "old": 0.25}) is None
    with pytest.raises(ValueError):
        standardize({"a": 1}, {"a": 2})
    for name in ["WHO2026", "WHO2021", "Segi", "Scandinavian", "European2013"]:
        assert sum(load_weights(name).values()) == pytest.approx(1)
    assert load_weights()["5_14"] == pytest.approx(14.24 / 100.02)
    with pytest.raises(ValueError):
        collapse_weights({"0-4": 1})


def synthetic_tables(con):
    population_rows, observation_rows = [], []
    for location, base_population, rate in [("AAA", 1000, 10), ("BBB", 9000, 100)]:
        for sex, sex_weight in [("both", 1), ("male", 0.6), ("female", 0.4)]:
            for age in AGES.values():
                pop = base_population * sex_weight * (7 if age == "all" else 1)
                population_rows.append((2023, location, sex, age, pop))
                for measure, factor in [("deaths", 1), ("dalys", 20)]:
                    for code in [0, 1]:
                        observation_rows.append(
                            (measure, 2023, location, sex, age, code, pop * rate * factor / 100000)
                        )
    populations = pd.DataFrame(
        population_rows, columns=["year", "location", "sex", "age_group", "population"]
    )
    observations = pd.DataFrame(
        observation_rows, columns=["measure", "year", "location", "sex", "age_group", "cause_code", "value"]
    )
    con.register("population_frame", populations)
    con.register("observation_frame", observations)
    con.execute("CREATE TABLE population AS SELECT * FROM population_frame")
    con.execute("CREATE TABLE observations AS SELECT * FROM observation_frame")
    references = con.execute("""SELECT 'CDR' AS measure,year,location,sex,age_group,cause_code,
            value/population*100000 AS value
        FROM observations JOIN population USING(year,location,sex,age_group)
        WHERE measure='deaths' AND age_group='all'""").df()
    references = pd.concat(
        [references, references.assign(measure="ASDR", value=references.value * 1.5)], ignore_index=True
    )
    con.register("reference_rates", references)
    mapping = pd.DataFrame(
        [
            {
                "code": 0,
                "name": "SYNTHETIC all causes",
                "level": 0,
                "parent": None,
                "organ": None,
                "relation": "none",
                "rationale": "Synthetic test root",
            },
            {
                "code": 1,
                "name": "SYNTHETIC heart cause",
                "level": 1,
                "parent": 0,
                "organ": "heart",
                "relation": "primary",
                "rationale": "Synthetic test leaf",
            },
        ]
    )
    return mapping


def test_shared_contract_fixture_and_synthetic_export(tmp_path, organs):
    validate_export(ROOT / "contracts/fixtures/data", organs)
    with duckdb.connect() as con:
        mapping = synthetic_tables(con)
        aggregate(con, mapping, organs)
        world = con.execute("""SELECT rate,ageStdRate FROM organ_burden WHERE location='WLD'
            AND sex='both' AND age_group='all' AND measure='deaths' AND organ='heart'""").fetchone()
        assert world == pytest.approx((91, 91))
        locations = pd.DataFrame({"location": ["AAA", "BBB"], "name": ["Synthetic A", "Synthetic B"]})
        result = export(con, mapping, locations, tmp_path, generated_at="2026-01-01T00:00:00Z")
        assert result["files"] == 6
        assert validate_export(tmp_path, organs) == 3
        before = {p.name: p.read_bytes() for p in tmp_path.rglob("*.json")}
        export(con, mapping, locations, tmp_path, generated_at="2026-01-01T00:00:00Z")
        assert before == {p.name: p.read_bytes() for p in tmp_path.rglob("*.json")}
        data = json.loads((tmp_path / "locations/WLD.json").read_text())
        assert data["values"]["deaths"]["2023"]["both"]["all"]["count"][organs.index("muscles")] == 0
        manifest = json.loads((tmp_path / "manifest.json").read_text())
        assert manifest["colorDomains"]["deaths.rate"] == [10, 100]
        data["values"]["deaths"]["2023"]["both"]["all"]["count"].pop()
        (tmp_path / "locations/WLD.json").write_text(json.dumps(data))
        with pytest.raises(AssertionError):
            validate_export(tmp_path, organs)


def test_missing_cause_propagates_to_organ_world_and_standardized(organs):
    with duckdb.connect() as con:
        mapping = synthetic_tables(con)
        con.execute(
            "UPDATE observations SET value=NULL WHERE cause_code=1 AND location='AAA' AND age_group='0_4'"
        )
        aggregate(con, mapping, organs)
        assert con.execute("""SELECT count,rate FROM organ_burden WHERE location='WLD' AND age_group='0_4'
            AND organ='heart' AND measure='deaths' AND sex='both'""").fetchone() == (None, None)
        assert con.execute("""SELECT ageStdRate FROM organ_burden WHERE location='WLD' AND age_group='all'
            AND organ='heart' AND measure='deaths' AND sex='both'""").fetchone() == (None,)


def test_country_deaths_use_published_sum_while_world_and_dalys_use_broad_bins(organs):
    with duckdb.connect() as con:
        mapping = synthetic_tables(con)
        aggregate(con, mapping, organs)
        values = con.execute("""SELECT measure,location,ageStdRate FROM organ_burden
            WHERE organ='heart' AND sex='both' AND age_group='all' ORDER BY measure,location""").fetchall()
        assert [(m, loc) for m, loc, _ in values] == [
            ("dalys", "AAA"),
            ("dalys", "BBB"),
            ("dalys", "WLD"),
            ("deaths", "AAA"),
            ("deaths", "BBB"),
            ("deaths", "WLD"),
        ]
        assert [value for _, _, value in values] == pytest.approx([200, 2000, 1820, 15, 150, 91])
        comparison = con.execute("""SELECT published_sum,broad_bin,diff,rel_diff
            FROM deaths_asdr_method_comparison WHERE location='AAA' AND sex='both' AND organ='heart'""").fetchone()
        assert comparison == pytest.approx((15.0, 10.0, -5.0, -1 / 3))
        empty = con.execute("""SELECT published_sum,broad_bin,diff,rel_diff
            FROM deaths_asdr_method_comparison WHERE location='AAA' AND sex='both' AND organ='muscles'""").fetchone()
        assert empty == (0.0, 0.0, 0.0, None)
        missing = con.execute("SELECT * FROM reference_rates").df()
        missing.loc[
            (missing.location == "AAA") & (missing.sex == "both") & (missing.cause_code == 1), "value"
        ] = None
        con.register("reference_rates", missing)
        aggregate(con, mapping, organs)
        assert con.execute("""SELECT ageStdRate FROM organ_burden WHERE measure='deaths'
            AND location='AAA' AND sex='both' AND age_group='all' AND organ='heart'""").fetchone() == (None,)


def test_reference_parser_keeps_cause_rates_in_per_100000_units(tiny_workbook):
    book = openpyxl.load_workbook(tiny_workbook)
    for name in list(book.sheetnames):
        if name not in ["Notes", "All ages"]:
            del book[name]
    sheet = book["All ages"]
    sheet.title = "ASDR 2023"
    sheet["A1"] = "SYNTHETIC age-standardized rate per 100 000"
    book.copy_worksheet(sheet).title = "ASDR 2020"
    book.save(tiny_workbook)
    rates = parse_reference(tiny_workbook, "ASDR")
    assert len(rates) == 18
    assert set(rates.cause_code) == {0, 10, 20}
    assert set(rates.year) == {2020, 2023}
    assert rates.value.max() == 1


def test_reference_validation_rejects_unexplained_partition_and_crude_errors():
    with duckdb.connect() as con:
        mapping = synthetic_tables(con)
        rates = con.execute("SELECT * FROM reference_rates").df()
        check = validate_reference_rates(con, mapping, rates)
        assert check["failures"] == []
        assert check["cause_level_crude_comparison"]["pairs"] == 12
        rates.loc[
            (rates.measure == "ASDR") & (rates.location == "AAA") & (rates.cause_code == 1), "value"
        ] += 1
        check = validate_reference_rates(con, mapping, rates)
        assert "published_ASDR_partition" in check["failures"]
        rates.loc[
            (rates.measure == "CDR") & (rates.location == "AAA") & (rates.cause_code == 1), "value"
        ] += 1
        assert "cause_level_crude_rates" in validate_reference_rates(con, mapping, rates)["failures"]
        absent = rates[~((rates.measure == "ASDR") & (rates.location == "AAA") & (rates.cause_code == 1))]
        with pytest.raises(ValueError, match="Missing published primary"):
            validate_reference_rates(con, mapping, absent)


def test_method_comparison_parquet_is_tidy_and_preserves_zero_denominators(tmp_path, organs):
    with duckdb.connect() as con:
        mapping = synthetic_tables(con)
        aggregate(con, mapping, organs)
        path = tmp_path / "synthetic_comparison.parquet"
        result = save_method_comparison(con, path)
        frame = con.from_parquet(str(path)).df()
        assert list(frame.columns) == [
            "location",
            "year",
            "sex",
            "organ",
            "published_sum",
            "broad_bin",
            "diff",
            "rel_diff",
        ]
        assert not frame.duplicated(["location", "year", "sex", "organ"]).any()
        assert len(frame) == result["rows"] == 2 * 3 * len(organs)
        assert frame.query("organ=='muscles'").rel_diff.isna().all()
        heart = next(row for row in result["by_organ"] if row["organ"] == "heart")
        assert heart["relative_pairs"] == 6
        assert heart["relative_difference_percent"]["median"] == pytest.approx(-100 / 3)


def test_sql_standardization_matches_pure_function_with_varying_age_rates(organs):
    rates = {age: float(index * 10) for index, age in enumerate(list(AGES.values())[1:], 1)}
    with duckdb.connect() as con:
        mapping = synthetic_tables(con)
        for age, rate in rates.items():
            con.execute(
                """UPDATE observations SET value=p.population*?/100000 FROM population p
                WHERE observations.year=p.year AND observations.location=p.location
                AND observations.sex=p.sex AND observations.age_group=p.age_group
                AND observations.age_group=?""",
                [rate, age],
            )
        aggregate(con, mapping, organs)
        value = con.execute("""SELECT ageStdRate FROM organ_burden WHERE organ='heart'
            AND age_group='all' AND location='WLD' AND measure='deaths' AND sex='both'""").fetchone()[0]
        assert value == pytest.approx(standardize(rates, load_weights()))
        assert (
            con.execute(
                "SELECT count(*) FROM organ_burden WHERE age_group!='all' AND ageStdRate IS NOT NULL"
            ).fetchone()[0]
            == 0
        )


@pytest.mark.parametrize(
    "filename", ["reconciliation_exceptions.csv", "reference_reconciliation_exceptions.csv"]
)
def test_reviewed_exceptions_require_the_exact_pinned_source_values(filename):
    path = CONFIG / filename
    snapshot = pd.read_csv(path).iloc[0]
    pair = snapshot.drop(labels=["rationale"]).to_dict()
    pair["expected"] = pair.pop("published_value")
    pair["actual"] = pair.pop("children_value")
    with duckdb.connect() as con:
        synthetic_tables(con)
        original = reviewed_exceptions(con, pd.DataFrame([pair]), "cause_code", path)
        assert original["unexplained_count"] == 0
        assert original["explained_counts"] == {"pinned_all_causes_defect": 1}
        pair["actual"] += 1
        assert reviewed_exceptions(con, pd.DataFrame([pair]), "cause_code", path)["unexplained_count"] == 1


def test_leishmaniasis_exception_requires_the_observed_source_identity():
    keys = {"measure": "deaths", "year": 2023, "location": "AAA", "sex": "both", "age_group": "all"}
    pairs = pd.DataFrame(
        [
            keys | {"parent": 10, "expected": 100.0, "actual": 97.5},
            keys | {"parent": 20, "expected": 50.0, "actual": 52.5},
        ]
    )
    with duckdb.connect() as con:
        synthetic_tables(con)
        con.execute("INSERT INTO observations VALUES ('deaths',2023,'AAA','both','all',260,2.5)")
        recognized = reviewed_exceptions(con, pairs, "parent")
        assert recognized["explained_counts"] == {"code_20_excludes_leishmaniasis": 2}
        assert recognized["unexplained_count"] == 0
        pairs.loc[0, "actual"] += 1
        assert reviewed_exceptions(con, pairs, "parent")["unexplained_count"] == 1


def test_no_nan():
    assert sig(float("nan")) is None
    assert number(None) is None
    assert sig(123456) == 123500
    assert json.dumps(sig(1000.0)) == "1000"
    with pytest.raises(ValueError):
        finite_tree({"test": [float("nan")]})
    with pytest.raises(ValueError):
        number(float("inf"))
    with pytest.raises(ValueError):
        sig(float("inf"))


def test_duckdb_parquet_roundtrip_and_reconciliation(tmp_path):
    with duckdb.connect() as con:
        mapping = synthetic_tables(con)
        frame = con.execute("SELECT * FROM observations").df()
        path = tmp_path / "synthetic.parquet"
        copy_parquet(con, frame, path)
        restored = con.from_parquet(str(path)).df()
        pd.testing.assert_frame_equal(frame, restored)
        checks = reconcile(con, mapping)
        assert checks["hierarchy_outside_tolerance"] == 0
        for name in ["sexes", "ages", "partition"]:
            assert checks[name]["outside_tolerance"] == 0
        con.execute("UPDATE observations SET value=value+1 WHERE cause_code=1 AND sex='male'")
        assert reconcile(con, mapping)["partition"]["outside_tolerance"] > 0


def test_schema_rejects_duplicate_negative_and_unknown_dimension():
    schema, _ = schemas()
    good = pd.DataFrame(
        [
            {
                "measure": "deaths",
                "year": 2023,
                "location": "AAA",
                "sex": "both",
                "age_group": "all",
                "cause_code": 0,
                "value": 1.0,
            }
        ]
    )
    schema.validate(good)
    for bad in [
        pd.concat([good, good], ignore_index=True),
        good.assign(value=-1),
        good.assign(sex="unknown"),
    ]:
        with pytest.raises(pa.errors.SchemaError):
            schema.validate(bad)
