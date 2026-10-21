"""Parse WHO GHE workbooks without trusting their cached worksheet dimensions."""

import math
import re
from dataclasses import dataclass

import openpyxl
import pandas as pd

from .causes import cause_from_row, hierarchy
from .constants import AGES, SEXES


@dataclass
class Parsed:
    observations: pd.DataFrame
    population: pd.DataFrame
    causes: pd.DataFrame
    locations: pd.DataFrame


def number(value, scale=1):
    if value is None or str(value).strip() in {"", "..", "...", "-"}:
        return None
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value):
        raise ValueError(f"Unexpected numeric cell: {value!r}")
    return float(value) * scale


def parse_sheet(ws, year, measure, age_group, scale=1000, all_cause_only=False):
    ws.reset_dimensions()
    rows = ws.iter_rows(values_only=True)
    header = None
    locations = None
    quality = None
    titles = []
    obs, pop, trees = [], [], {}
    populations_seen = set()
    for row in rows:
        row = tuple(row)
        if len(row) < 2:
            if header is None:
                titles.extend(str(v) for v in row if v is not None)
            continue
        if row[0] == "Sex" and row[1] == "GHE code":
            if header is not None:
                raise ValueError(f"Repeated header in {ws.title}")
            header = row
            title = " ".join(titles)
            if scale == 1000 and "('000)" not in title:
                raise ValueError("Expected estimates in thousands in the workbook title")
            if scale == 1 and "per 100 000" not in title:
                raise ValueError("Expected published reference rates per 100,000")
            continue
        if "ISO-3 Code" in row:
            if header is None or locations is not None:
                raise ValueError("ISO row must follow the country header")
            locations = [
                (i, v, str(header[i]).strip()) for i, v in enumerate(row) if i >= 7 and v is not None
            ]
            if not locations or any(not re.fullmatch(r"[A-Z]{3}", v) for _, v, _ in locations):
                raise ValueError("Invalid country columns")
            if len({v for _, v, _ in locations}) != len(locations):
                raise ValueError("Duplicate ISO code")
            continue
        if locations is None:
            titles.extend(str(v) for v in row if v is not None)
            continue
        if quality is None:
            quality = {iso: number(row[i]) for i, iso, _ in locations}
            if any(v not in {1, 2, 3, 4} for v in quality.values()):
                raise ValueError("Expected WHO evidence-category row after ISO codes")
            continue
        if row[0] not in SEXES:
            if row[0] is not None and isinstance(row[1], (int, float)):
                raise ValueError(f"Unrecognized sex: {row[0]}")
            continue
        sex = SEXES[row[0]]
        population_row = any(str(v).startswith("Population ('000)") for v in row[2:7])
        if population_row:
            if sex in populations_seen:
                raise ValueError(f"Repeated population row: {sex}")
            populations_seen.add(sex)
            pop.extend((year, iso, sex, age_group, number(row[i], 1000)) for i, iso, _ in locations)
        elif isinstance(row[1], (int, float)):
            if sex not in populations_seen:
                raise ValueError("Cause block precedes population")
            cause = cause_from_row(row)
            trees.setdefault(sex, []).append(cause)
            if not all_cause_only or cause["code"] == 0:
                obs.extend(
                    (measure, year, iso, sex, age_group, cause["code"], number(row[i], scale))
                    for i, iso, _ in locations
                )
        elif any(v is not None for v in row[2:7]):
            raise ValueError(f"Unexpected data row in {ws.title}: {row[:7]}")
    if populations_seen != set(SEXES.values()) or set(trees) != set(SEXES.values()):
        raise ValueError(f"Incomplete sex blocks in {ws.title}")
    tree = hierarchy(trees["both"])
    if any(hierarchy(v) != tree for v in trees.values()):
        raise ValueError("Cause definitions differ between sexes")
    return Parsed(
        pd.DataFrame(obs, columns=["measure", "year", "location", "sex", "age_group", "cause_code", "value"]),
        pd.DataFrame(pop, columns=["year", "location", "sex", "age_group", "population"]),
        pd.DataFrame(tree),
        pd.DataFrame(
            [{"location": iso, "name": name, "data_quality": quality[iso]} for _, iso, name in locations]
        ),
    )


def parse_workbook(path, year, measure):
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if set(book.sheetnames) != {"Notes", *AGES}:
            raise ValueError(f"Unexpected sheets: {book.sheetnames}")
        parts = [parse_sheet(book[sheet], year, measure, age) for sheet, age in AGES.items()]
        for part in parts[1:]:
            if not part.causes.equals(parts[0].causes) or not part.locations.equals(parts[0].locations):
                raise ValueError("Cause/country definitions differ between age sheets")
        return Parsed(
            pd.concat([p.observations for p in parts], ignore_index=True),
            pd.concat([p.population for p in parts], ignore_index=True),
            parts[0].causes,
            parts[0].locations,
        )
    finally:
        book.close()


def parse_reference(path, prefix):
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheets = [s for s in book.sheetnames if s != "Notes"]
        if not sheets or any(not re.fullmatch(rf"{prefix} \d{{4}}", s) for s in sheets):
            raise ValueError(f"Unexpected reference sheets: {book.sheetnames}")
        years = sorted(int(s.split()[-1]) for s in sheets)
        parts = [parse_sheet(book[f"{prefix} {year}"], year, prefix, "all", scale=1) for year in years]
        for part in parts[1:]:
            if not part.causes.equals(parts[0].causes) or not part.locations.equals(parts[0].locations):
                raise ValueError("Cause/country definitions differ between reference years")
        return pd.concat([part.observations for part in parts], ignore_index=True)
    finally:
        book.close()
